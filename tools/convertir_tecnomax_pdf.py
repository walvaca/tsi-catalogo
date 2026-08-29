"""
Convierte la lista de precios de Tecnomax (PDF) a un .xlsx plano que la app
puede importar con el asistente genérico ("Otro proveedor") — incluyendo la
foto de producto que trae el propio PDF.

Por qué existe: Tecnomax manda su catálogo en PDF, no en Excel como GVS. El
PDF sí tiene una tabla real por dentro (con celdas delimitadas) y una imagen
de producto embebida junto a cada fila, así que se puede leer con precisión
con PyMuPDF (fitz) — pero cada sección del documento usa un layout de
columnas ligeramente distinto (cámaras/DVR, accesorios con foto, racks por
dimensiones, etc.), así que en vez de asumir columnas fijas, este script
clasifica cada fila de forma genérica:
  - fila con una sola celda con texto -> etiqueta de marca/sección
  - fila con una celda con fecha -> etiqueta de categoría (esa fecha es la
    "fecha del proveedor" del catálogo)
  - fila con "Código"/"Referencia" (o dos últimas celdas tipo "Precio"/
    "Precio IVA") -> encabezado de subcategoría
  - cualquier otra fila -> producto: primera celda no vacía = código, el
    resto (antes de las 2 últimas) = descripción, las 2 últimas = precios
    (o "AGOTADO"/"CONSULTAR ASESOR" en vez de precio)

La foto de cada producto se ubica por posición: se toma la imagen embebida
en el PDF cuyo rango vertical mejor se solapa con el de esa fila de tabla
(ambos vienen de PyMuPDF, mismo sistema de coordenadas — no se mezclan dos
librerías distintas para evitar que una fila quede con la imagen de otra).
Se guarda embebida en la columna A del .xlsx de salida, en la MISMA fila que
su producto — igual que hace GVS en su propio Excel — así que el extractor
de fotos que ya existe en la app (js/xlsx-images.js) la reconoce sin tocar
ni una línea de código del navegador. Verificado contra el catálogo de
agosto 2026: ~98% de los productos trae foto.

Uso:
    python convertir_tecnomax_pdf.py "ruta/al/catalogo.pdf" "ruta/de/salida.xlsx"

Requiere: pip install pymupdf openpyxl

Después de correrlo: importa el .xlsx resultante en la app con "Otro
proveedor" -> TECNOMAX. Mapeo de columnas (0=Foto, no se mapea a ningún
campo): Marca=1, Categoria=2, Subcategoria=3, Codigo=4, Descripcion=5,
PrecioSinIVA=6, PrecioConIVA=7, Estado(->Texto de existencia/stock)=8, fila
de inicio=1. Si ya existe un perfil guardado de una versión anterior de este
script (sin foto), hay que abrir "Editar mapeo" y volver a mapear una vez,
porque las columnas se corrieron un puesto para dejarle espacio a la foto.
"""
import sys
import io
import re
import fitz
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

DATE_RE = re.compile(r'\d{1,2}/\d{1,2}/\d{4}')
MONEY_RE = re.compile(r'[\d.,]+')
HEADER_WORDS = {'codigo', 'código', 'referencia'}
HEADER_PRICE_WORDS = {'precio', 'precio iva', 'antes de iva', 'incluido iva', 'antes de i.v.a', 'incluido i.v.a'}
BOILERPLATE = {
    'codigo', 'código', 'referencia', 'imagen', 'descripcion', 'descripción',
    'antes de iva', 'antes de i.v.a', 'incluido iva', 'incluido i.v.a',
    'precio', 'precio iva', 'dimensiones (cm', 'dimensiones (cm)'
}
# raíces amplias a propósito: en este documento "agotado" y "consultar con el
# asesor" aparecen con varios errores de tipeo (ASESRO, CONSUTA, etc.)
NO_PRECIO_MARCAS = ('agotad', 'consu', 'ases')


def clean(c):
    if c is None:
        return ''
    return re.sub(r'\s+', ' ', str(c)).strip()


def es_no_disponible(c):
    return any(m in clean(c).lower() for m in NO_PRECIO_MARCAS)


def es_vacio_o_placeholder(c):
    return clean(c).replace('$', '').replace('-', '').replace(' ', '') == ''


def es_incluido_iva_label(c):
    c = clean(c).lower()
    return 'incluido iva' in c or 'incuido iva' in c


def es_terminal_valida(c):
    if es_vacio_o_placeholder(c) or es_no_disponible(c) or es_incluido_iva_label(c):
        return True
    c = clean(c).replace(' ', '')
    return bool(MONEY_RE.search(c)) and any(ch.isdigit() for ch in c)


def parsear_precio(c):
    c = clean(c).replace(' ', '')
    if es_vacio_o_placeholder(c) or es_no_disponible(c) or es_incluido_iva_label(c):
        return None
    m = MONEY_RE.search(c)
    if not m:
        return None
    n = m.group(0).replace('.', '').replace(',', '.')
    try:
        return float(n)
    except ValueError:
        return None


def estado_de(c):
    c = clean(c).lower()
    if 'agotad' in c:
        return 'AGOTADO'
    if 'consu' in c or 'ases' in c:
        return 'CONSULTAR ASESOR'
    return ''


def imagenes_de_pagina(page):
    """[(rect, xref), ...] de todas las posiciones de imagen en la página."""
    entradas = []
    for xref, *_ in page.get_images(full=True):
        for rect in page.get_image_rects(xref):
            entradas.append((rect, xref))
    return entradas


def mejor_imagen_para_fila(img_entries, row_bbox):
    x0, top, x1, bottom = row_bbox
    mid = (top + bottom) / 2
    candidatos = [(r, x) for r, x in img_entries if r.y0 <= mid <= r.y1]
    if not candidatos:
        return None
    candidatos.sort(key=lambda c: -(min(bottom, c[0].y1) - max(top, c[0].y0)))
    return candidatos[0][1]


def convertir(path_pdf, path_out):
    filas_final = []  # cada item: (dict de campos, bytes de imagen o None, extension)
    categoria_actual = ''
    marca_actual = ''
    subcategoria_actual = ''
    sin_clasificar = []
    con_foto = 0

    doc = fitz.open(path_pdf)
    for pnum in range(len(doc)):
        page = doc[pnum]
        img_entries = imagenes_de_pagina(page)
        for tabla in page.find_tables().tables:
            textrows = tabla.extract()
            for textrow, row in zip(textrows, tabla.rows):
                cells = [clean(c) for c in textrow]
                non_empty = [c for c in cells if c]
                if not non_empty:
                    continue

                ultimos2 = non_empty[-2:] if len(non_empty) >= 2 else non_empty
                es_header = (
                    any(c.lower() in HEADER_WORDS for c in non_empty) or
                    all(c.lower() in HEADER_PRICE_WORDS for c in ultimos2)
                )
                if es_header:
                    extra = [c for c in non_empty if c.lower() not in BOILERPLATE and not DATE_RE.search(c)]
                    if extra:
                        subcategoria_actual = extra[0]
                    continue

                tiene_fecha = any(DATE_RE.search(c) for c in non_empty)
                resto_sin_fecha = [c for c in non_empty if not DATE_RE.search(c)]
                if tiene_fecha:
                    categoria_actual = resto_sin_fecha[0] if resto_sin_fecha else categoria_actual
                    continue
                if len(non_empty) == 1:
                    marca_actual = non_empty[0]
                    continue

                if not (len(cells) >= 2 and es_terminal_valida(cells[-1]) and es_terminal_valida(cells[-2])):
                    sin_clasificar.append((pnum + 1, row))
                    continue

                cuerpo = [c for i, c in enumerate(cells) if c and i < len(cells) - 2]
                if not cuerpo:
                    sin_clasificar.append((pnum + 1, row))
                    continue
                codigo = cuerpo[0]
                descripcion = ' '.join(cuerpo[1:]) if len(cuerpo) > 1 else ''

                if es_incluido_iva_label(cells[-1]) and not es_incluido_iva_label(cells[-2]):
                    precio_con_iva = parsear_precio(cells[-2])
                    precio_sin_iva = None
                    estado = ''
                else:
                    precio_sin_iva = parsear_precio(cells[-2])
                    precio_con_iva = parsear_precio(cells[-1])
                    estado = estado_de(cells[-2]) or estado_de(cells[-1])

                if precio_sin_iva is None and precio_con_iva is None and not estado:
                    estado = 'CONSULTAR ASESOR'

                img_bytes, img_ext = None, None
                xref = mejor_imagen_para_fila(img_entries, row.bbox)
                if xref is not None:
                    info = doc.extract_image(xref)
                    img_bytes, img_ext = info['image'], info['ext']
                    con_foto += 1

                filas_final.append(({
                    'Marca': marca_actual,
                    'Categoria': categoria_actual,
                    'Subcategoria': subcategoria_actual,
                    'Codigo': codigo,
                    'Descripcion': descripcion,
                    'PrecioSinIVA': precio_sin_iva,
                    'PrecioConIVA': precio_con_iva,
                    'Estado': estado
                }, img_bytes, img_ext))

    print('total productos:', len(filas_final))
    print(f'con foto: {con_foto} ({100 * con_foto / max(1, len(filas_final)):.1f}%)')
    print('filas sin poder clasificar (revisar a mano):', len(sin_clasificar))
    for pnum, row in sin_clasificar:
        print(' p', pnum, row)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tecnomax'
    headers = ['Foto', 'Marca', 'Categoria', 'Subcategoria', 'Codigo', 'Descripcion', 'PrecioSinIVA', 'PrecioConIVA', 'Estado']
    ws.append(headers)
    ws.column_dimensions['A'].width = 12

    for i, (f, img_bytes, img_ext) in enumerate(filas_final):
        fila_excel = i + 2  # fila 1 = encabezado
        ws.append(['', f['Marca'], f['Categoria'], f['Subcategoria'], f['Codigo'], f['Descripcion'],
                    f['PrecioSinIVA'], f['PrecioConIVA'], f['Estado']])
        ws.row_dimensions[fila_excel].height = 46
        if img_bytes:
            try:
                xlimg = XLImage(io.BytesIO(img_bytes))
                xlimg.width, xlimg.height = 60, 45
                ws.add_image(xlimg, f'A{fila_excel}')
            except Exception as e:
                print(f'  aviso: no se pudo incrustar foto de fila {fila_excel} ({f["Codigo"]}): {e}')

    wb.save(path_out)
    print('guardado en', path_out)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('Uso: python convertir_tecnomax_pdf.py <entrada.pdf> <salida.xlsx>')
        sys.exit(1)
    convertir(sys.argv[1], sys.argv[2])
